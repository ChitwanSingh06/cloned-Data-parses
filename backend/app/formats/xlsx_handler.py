"""XLSX -> canonical Document extraction.

Each worksheet is a logical page. Spreadsheet content stays structured as canonical
``table`` blocks rather than being flattened into prose. Physical PDF coordinates
are not invented; provenance records worksheet/range/cell coordinates instead.
"""
from __future__ import annotations

import datetime as _dt
import time
import uuid
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.core.config import PROCESSED_DIR
from app.models.block import Block, BlockType
from app.models.document import Document, Page
from app.models.table import Cell, TableData
from app.output.json_builder import build_json
from app.output.markdown_builder import build_markdown
from app.pipeline.confidence import add_confidence
from app.pipeline.failsafe import final_status, make_error
from app.pipeline.provenance import add_provenance
from app.pipeline.timing import StageTimer, compute_metrics
from app.summary.summarizer import build_summary
from app.utils.platform_utils import peak_memory_mb


def _value(value: Any) -> Any:
    """Return an XLSX cell value that is safe in the canonical JSON model."""
    if value is None:
        return ""
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return value


def _used_range(ws) -> tuple[int, int, int, int] | None:
    """Find the smallest rectangle containing actual cell values, not formatting."""
    min_row = min_col = None
    max_row = max_col = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            r, c = cell.row, cell.column
            min_row = r if min_row is None else min(min_row, r)
            min_col = c if min_col is None else min(min_col, c)
            max_row = max(max_row, r)
            max_col = max(max_col, c)
    if min_row is None:
        return None
    return min_row, min_col, max_row, max_col


def _looks_like_header(values: list[list[Any]]) -> bool:
    if len(values) < 2 or not values[0]:
        return False
    first = values[0]
    nonempty = [v for v in first if str(v).strip()]
    if len(nonempty) < 1:
        return False
    # Conservative: a row with labels and at least one additional row is useful as
    # Markdown headers. The original values remain in canonical cells regardless.
    return all(not isinstance(v, (int, float, bool)) for v in nonempty)


def _worksheet_block(ws, sheet_index: int, block_id: str) -> tuple[Block | None, int, int, dict[str, Any]]:
    used = _used_range(ws)
    if used is None:
        return None, 0, 0, {}
    min_row, min_col, max_row, max_col = used
    values: list[list[Any]] = []
    cells: list[Cell] = []
    cell_provenance: list[dict[str, Any]] = []

    for r in range(min_row, max_row + 1):
        row_values: list[Any] = []
        for c in range(min_col, max_col + 1):
            cell = ws.cell(r, c)
            value = _value(cell.value)
            row_values.append(value)
            cells.append(Cell(row=r - min_row, col=c - min_col, text=str(value) if value != "" else ""))
            cell_provenance.append({
                "worksheet": ws.title,
                "worksheet_index": sheet_index,
                "cell": cell.coordinate,
                "row": r,
                "column": c,
                "column_letter": get_column_letter(c),
            })
        values.append(row_values)

    header = _looks_like_header(values)
    headers = [values[0]] if header else []
    body = values[1:] if header else values
    table_names = []
    try:
        table_names = [name for name in ws.tables.keys()]
    except Exception:
        pass

    excel_tables = []
    for name in table_names:
        try:
            excel_tables.append({"name": name, "ref": ws.tables[name].ref})
        except Exception:
            pass

    data = TableData(
        n_rows=len(values),
        n_cols=max_col - min_col + 1,
        header_rows=1 if header else 0,
        headers=headers,
        rows=body,
        cells=cells,
        page_span=[sheet_index],
    )
    meta = {
        "source_type": "xlsx",
        "worksheet": ws.title,
        "worksheet_index": sheet_index,
        "range": f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}",
        "start_row": min_row,
        "start_column": min_col,
        "end_row": max_row,
        "end_column": max_col,
        "cell_provenance": cell_provenance,
        "header_row_inferred": header,
        "excel_tables": excel_tables,
        "region_id": block_id,
        "sources": [{
            "page": sheet_index,
            "bbox": None,
            "region_id": block_id,
            "source_type": "xlsx",
            "worksheet": ws.title,
            "worksheet_index": sheet_index,
            "range": f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}",
        }],
    }
    signals = {
        "structural_extraction": 0.98,
        "worksheet_structure": 0.98,
        "cell_values": 0.98,
    }
    if not values:
        signals["cell_values"] = 0.4
    block = Block(
        id=block_id,
        type=BlockType.table,
        content=data.model_dump(),
        page=sheet_index,
        bbox=None,
        extractor="openpyxl:worksheet",
        meta=meta,
        signals=signals,
    )
    return block, len(values), sum(len(str(v)) for row in values for v in row if v != ""), meta


def _sheet_heading(ws, sheet_index: int, block_id: str) -> Block:
    return Block(
        id=block_id,
        type=BlockType.heading,
        content=ws.title,
        page=sheet_index,
        bbox=None,
        extractor="openpyxl:worksheet",
        level=1,
        meta={
            "source_type": "xlsx",
            "worksheet": ws.title,
            "worksheet_index": sheet_index,
            "region_id": block_id,
            "sources": [{
                "page": sheet_index,
                "bbox": None,
                "region_id": block_id,
                "source_type": "xlsx",
                "worksheet": ws.title,
                "worksheet_index": sheet_index,
            }],
        },
        signals={"structural_extraction": 0.98, "worksheet_name": 0.99},
    )


def parse_xlsx(path: str, document_id: str | None = None, filename: str | None = None,
               out_root: Path | None = None) -> Document:
    """Parse XLSX into the existing canonical Document model and outputs."""
    t0 = time.perf_counter()
    timer = StageTimer()
    document_id = document_id or uuid.uuid4().hex[:12]
    filename = filename or Path(path).name
    out_dir = (out_root or PROCESSED_DIR) / document_id
    out_dir.mkdir(parents=True, exist_ok=True)
    try:
        with timer.stage("ingestion"):
            wb = load_workbook(filename=path, data_only=False, read_only=False)
            sheet_count = len(wb.worksheets)

        pages: list[Page] = []
        blocks: list[Block] = []
        errors: list[dict[str, Any]] = []
        counter = 0

        with timer.stage("layout_analysis"):
            for sheet_index, ws in enumerate(wb.worksheets, start=1):
                page_blocks: list[str] = []
                sheet_chars = 0
                counter += 1
                heading = _sheet_heading(ws, sheet_index, f"X{counter:04d}")
                blocks.append(heading)
                page_blocks.append(heading.id)

                counter += 1
                table, n_rows, chars, _ = _worksheet_block(ws, sheet_index, f"X{counter:04d}")
                if table is not None:
                    blocks.append(table)
                    page_blocks.append(table.id)
                    sheet_chars += chars
                else:
                    # A worksheet with no values is still represented as a logical page.
                    # The sheet heading preserves the worksheet identity without inventing data.
                    pass

                pages.append(Page(
                    page_number=sheet_index,
                    width=float(ws.max_column or 0),
                    height=float(ws.max_row or 0),
                    is_scanned=False,
                    text_chars=sheet_chars + len(ws.title),
                    source="digital",
                    blocks=page_blocks,
                ))

        wb.close()
        doc = Document(
            document_id=document_id,
            filename=filename,
            format="xlsx",
            page_count=sheet_count,
            pages=pages,
            blocks=blocks,
            errors=errors,
        )

        with timer.stage("assembly"):
            # Worksheet/row/cell order is already the canonical structural order.
            pass
        with timer.stage("validation"):
            add_confidence(doc)
            add_provenance(doc)
        doc.status = final_status(doc.errors, len(doc.blocks))

        with timer.stage("summary"):
            try:
                doc.summary = build_summary(doc)
            except Exception as exc:
                doc.summary = {}
                doc.errors.append(make_error("PARSING_FAILED", f"Summary generation failed: {exc}", severity="warning"))

        elapsed = time.perf_counter() - t0
        doc.processing_time = round(elapsed, 3)
        doc.timing = {
            **compute_metrics(elapsed, sheet_count),
            "logical_units_processed": sheet_count,
            "unit_label": "worksheets",
            "stages": timer.as_dict(),
        }
        counts = Counter(b.type.value for b in doc.blocks)
        doc.stats = {
            "pages_processed": sheet_count,
            "logical_units_processed": sheet_count,
            "unit_label": "worksheets",
            "seconds_per_page": doc.timing["seconds_per_page"],
            "peak_memory_mb": peak_memory_mb(),
            "scanned_pages": 0,
            "digital_pages": sheet_count,
            "blocks_by_type": dict(counts),
            "review_required": sum(b.status.value == "REVIEW_REQUIRED" for b in doc.blocks),
            "confidence_levels": dict(Counter(b.confidence_level.value for b in doc.blocks)),
            "mean_confidence": round(sum(b.confidence for b in doc.blocks) / len(doc.blocks), 3) if doc.blocks else 0.0,
            "errors": sum(e.get("severity") == "error" for e in doc.errors),
            "warnings": sum(e.get("severity") == "warning" for e in doc.errors),
        }
        payload = doc.model_dump(mode="json")
        (out_dir / "document.md").write_text(build_markdown(payload), encoding="utf-8", newline="\n")
        payload = doc.model_dump(mode="json")
        (out_dir / "document.json").write_text(build_json(payload), encoding="utf-8", newline="\n")
        return doc
    except Exception as exc:
        doc = Document(
            document_id=document_id,
            filename=filename,
            format="xlsx",
            status="FAILED",
            errors=[make_error("CORRUPT_XLSX" if isinstance(exc, (ValueError, OSError, RuntimeError, KeyError)) else "PARSING_FAILED", str(exc))],
        )
        return doc


def handle_xlsx(path: str, **kw) -> Document:
    return parse_xlsx(path, **kw)
