"""Embedded pictures are detected as `figure` blocks in every supported format."""
import io
import random

import pytest
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw

from app.main import app

client = TestClient(app)


@pytest.fixture(scope="module")
def png(tmp_path_factory):
    random.seed(1)
    im = Image.new("RGB", (400, 300))
    d = ImageDraw.Draw(im)
    for _ in range(60):
        x, y = random.randint(0, 350), random.randint(0, 250)
        d.ellipse([x, y, x + random.randint(20, 50), y + random.randint(20, 50)],
                  fill=tuple(random.randint(0, 255) for _ in range(3)))
    path = tmp_path_factory.mktemp("img") / "photo.png"
    im.save(path)
    return str(path)


def _parse(name: str, data: bytes):
    up = client.post("/api/upload", files={"file": (name, data)}).json()
    assert client.post(f"/api/parse?file_id={up['file_id']}&wait=true").status_code == 200
    return up["file_id"], client.get(f"/api/result/{up['file_id']}").json()


def _figures(res):
    return [b for b in res["blocks"] if b["type"] == "figure"]


def _assert_image_served(fid, fig):
    path = fig["meta"]["image_path"]
    r = client.get(f"/api/documents/{fid}/figures/{path.split('/')[-1]}")
    assert r.status_code == 200 and r.content[:4] == b"\x89PNG"


def test_pdf(png):
    import pymupdf
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 72), "Report with a picture below. This page has enough text to count as digital text.")
    p.insert_image(pymupdf.Rect(72, 120, 372, 345), filename=png)
    fid, res = _parse("t.pdf", doc.tobytes())
    figs = _figures(res)
    assert len(figs) == 1 and figs[0]["bbox"]
    _assert_image_served(fid, figs[0])


def test_docx(png):
    import docx
    from docx.shared import Inches
    d = docx.Document()
    d.add_paragraph("Before")
    d.add_picture(png, width=Inches(3))
    buf = io.BytesIO()
    d.save(buf)
    fid, res = _parse("t.docx", buf.getvalue())
    figs = _figures(res)
    assert len(figs) == 1
    _assert_image_served(fid, figs[0])


def test_pptx(png):
    from pptx import Presentation
    from pptx.util import Inches
    pr = Presentation()
    s = pr.slides.add_slide(pr.slide_layouts[5])
    s.shapes.add_picture(png, Inches(1), Inches(2), Inches(4))
    buf = io.BytesIO()
    pr.save(buf)
    fid, res = _parse("t.pptx", buf.getvalue())
    figs = _figures(res)
    assert len(figs) == 1
    _assert_image_served(fid, figs[0])


def test_xlsx_two_sheets(png):
    import openpyxl
    from openpyxl.drawing.image import Image as XI
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "a", "b"
    ws.add_image(XI(png), "D3")
    ws2 = wb.create_sheet("Second")
    ws2["A1"] = "x"
    ws2.add_image(XI(png), "B2")
    buf = io.BytesIO()
    wb.save(buf)
    fid, res = _parse("t.xlsx", buf.getvalue())
    figs = _figures(res)
    assert len(figs) == 2
    assert [f["page"] for f in figs] == [1, 2]
    assert figs[0]["meta"]["anchor_cell"] == "D3" and figs[1]["meta"]["anchor_cell"] == "B2"
    assert len({f["id"] for f in res["blocks"]}) == len(res["blocks"])  # block ids stay unique
    for f in figs:
        _assert_image_served(fid, f)
    assert res["status"] == "SUCCESS"


def test_xlsx_without_images_unchanged():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active["A1"] = "only text"
    buf = io.BytesIO()
    wb.save(buf)
    _, res = _parse("plain.xlsx", buf.getvalue())
    assert _figures(res) == [] and res["status"] == "SUCCESS"
