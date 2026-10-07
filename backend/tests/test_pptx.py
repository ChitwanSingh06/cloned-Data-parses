from pathlib import Path

from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from app.formats.pptx_handler import parse_pptx
from app.main import app
from app.models.block import BlockType

client = TestClient(app)


def _fixture(tmp_path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title = slide.shapes.title
    title.text = "PPTX Test Heading"
    box = slide.shapes.add_textbox(Inches(1), Inches(1.7), Inches(7), Inches(1.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = "First slide paragraph extracted from the test presentation."
    p.alignment = PP_ALIGN.LEFT
    for text in ("Bullet one", "Bullet two"):
        q = tf.add_paragraph()
        q.text = text
        q.level = 1
        q._p.get_or_add_pPr().append(__import__('pptx').oxml.parse_xml('<a:buChar xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" char="•"/>'))
    ordered_box = slide.shapes.add_textbox(Inches(6), Inches(1.7), Inches(3), Inches(1.2))
    ordered_tf = ordered_box.text_frame
    for i, text in enumerate(("Step one", "Step two")):
        q = ordered_tf.paragraphs[0] if i == 0 else ordered_tf.add_paragraph()
        q.text = text
        q._p.get_or_add_pPr().append(__import__('pptx').oxml.parse_xml('<a:buAutoNum xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" type="arabicPeriod"/>'))
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(3.4), Inches(5), Inches(1.5)).table
    table.cell(0, 0).text = "Column A"
    table.cell(0, 1).text = "Column B"
    table.cell(1, 0).text = "A1"
    table.cell(1, 1).text = "B1"
    table.cell(2, 0).text = "A2"
    table.cell(2, 1).text = "B2"
    image = tmp_path / "fixture.png"
    Image.new("RGB", (32, 24), "white").save(image)
    slide.shapes.add_picture(str(image), Inches(7), Inches(3.3), width=Inches(1))

    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Second Slide"
    slide2.placeholders[1].text = "Second slide paragraph."
    out = tmp_path / "fixture.pptx"
    prs.save(out)
    return out


def test_pptx_parsing_canonical_output(tmp_path):
    path = _fixture(tmp_path)
    doc = parse_pptx(str(path), document_id="pptx12345678")
    assert doc.format == "pptx"
    assert doc.status == "SUCCESS"
    assert doc.page_count == 2
    assert len(doc.pages) == 2
    assert doc.pages[0].width > 0 and doc.pages[0].height > 0
    assert [b.type for b in doc.blocks[:2]] == [BlockType.heading, BlockType.paragraph]
    assert any(b.type == BlockType.list and not b.content["ordered"] for b in doc.blocks)
    assert any(b.type == BlockType.list and b.content["ordered"] for b in doc.blocks)
    table = next(b for b in doc.blocks if b.type == BlockType.table)
    assert table.content["n_rows"] == 3 and table.content["n_cols"] == 2
    assert table.content["cells"][0]["text"] == "Column A"
    fig = next(b for b in doc.blocks if b.type == BlockType.figure)
    assert fig.meta["source_type"] == "pptx" and fig.meta["image_index"] == 1
    assert fig.meta["image_path"].endswith(".png")
    heading = next(b for b in doc.blocks if b.type == BlockType.heading)
    assert heading.page == 1
    assert heading.bbox and len(heading.bbox) == 4
    assert heading.provenance and heading.provenance.sources[0]["slide_number"] == 1
    from app.core.config import PROCESSED_DIR
    out = PROCESSED_DIR / "pptx12345678"
    assert (out / "document.json").exists()
    assert (out / "document.md").exists()


def test_pptx_upload_parse_downloads(tmp_path):
    path = _fixture(tmp_path)
    with path.open("rb") as fh:
        up = client.post("/api/upload", files={"file": ("fixture.pptx", fh, "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
    assert up.status_code == 200
    fid = up.json()["file_id"]
    st = client.post(f"/api/parse?file_id={fid}&wait=true").json()
    assert st["status"] == "SUCCESS"
    result = client.get(f"/api/result/{fid}").json()
    assert result["format"] == "pptx" and result["page_count"] == 2
    assert client.get(f"/api/download/{fid}/json").status_code == 200
    md = client.get(f"/api/download/{fid}/markdown")
    assert "PPTX Test Heading" in md.text and "Column A" in md.text
    assert client.get(f"/api/documents/{fid}/page/1.png").status_code == 404
