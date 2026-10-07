from pathlib import Path

from docx import Document as DocxDocument
from docx.shared import Inches
from PIL import Image

from app.formats.docx_handler import parse_docx
from app.models.block import BlockType
from app.main import app
from fastapi.testclient import TestClient

client = TestClient(app)


def _fixture(tmp_path: Path) -> Path:
    doc = DocxDocument()
    doc.add_heading("DOCX Test Heading", level=1)
    doc.add_paragraph("First paragraph extracted from the test document.")
    doc.add_paragraph("Second paragraph with structural provenance.")
    doc.add_paragraph("Bullet one", style="List Bullet")
    doc.add_paragraph("Bullet two", style="List Bullet")
    doc.add_paragraph("Step one", style="List Number")
    doc.add_paragraph("Step two", style="List Number")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Column A"
    table.cell(0, 1).text = "Column B"
    table.cell(1, 0).text = "A1"
    table.cell(1, 1).text = "B1"
    table.cell(2, 0).text = "A2"
    table.cell(2, 1).text = "B2"
    image = tmp_path / "fixture.png"
    Image.new("RGB", (32, 24), "white").save(image)
    doc.add_paragraph().add_run().add_picture(str(image), width=Inches(0.4))
    out = tmp_path / "fixture.docx"
    doc.save(out)
    return out


def test_docx_parsing_canonical_output(tmp_path):
    path = _fixture(tmp_path)
    doc = parse_docx(str(path), document_id="abc123abc123")
    assert doc.format == "docx"
    assert doc.status == "SUCCESS"
    assert doc.blocks
    assert [b.type for b in doc.blocks[:3]] == [BlockType.heading, BlockType.paragraph, BlockType.paragraph]
    assert any(b.type == BlockType.list and not b.content["ordered"] for b in doc.blocks)
    assert any(b.type == BlockType.list and b.content["ordered"] for b in doc.blocks)
    table = next(b for b in doc.blocks if b.type == BlockType.table)
    assert table.content["n_rows"] == 3 and table.content["n_cols"] == 2
    assert table.content["cells"][0]["text"] == "Column A"
    fig = next(b for b in doc.blocks if b.type == BlockType.figure)
    assert fig.meta["source_type"] == "docx" and fig.meta["image_index"] == 1
    para = next(b for b in doc.blocks if b.type == BlockType.paragraph)
    assert para.meta["source_type"] == "docx" and "paragraph_index" in para.meta
    assert para.bbox is None
    assert para.provenance and para.provenance.bbox is None
    assert para.provenance.sources[0]["source_type"] == "docx"
    from app.core.config import PROCESSED_DIR
    out = PROCESSED_DIR / "abc123abc123"
    assert (out / "document.json").exists()
    assert (out / "document.md").exists()


def test_docx_upload_parse_downloads(tmp_path, monkeypatch):
    path = _fixture(tmp_path)
    with path.open("rb") as fh:
        up = client.post("/api/upload", files={"file": ("fixture.docx", fh, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert up.status_code == 200
    fid = up.json()["file_id"]
    st = client.post(f"/api/parse?file_id={fid}&wait=true").json()
    assert st["status"] == "SUCCESS"
    result = client.get(f"/api/result/{fid}").json()
    assert result["format"] == "docx"
    assert client.get(f"/api/download/{fid}/json").status_code == 200
    md = client.get(f"/api/download/{fid}/markdown")
    assert "DOCX Test Heading" in md.text and "Column A" in md.text
    assert client.get(f"/api/documents/{fid}/page/1.png").status_code == 404


def test_unsupported_format_still_structured(tmp_path):
    path = tmp_path / "sample.txt"
    path.write_text("not a supported document", encoding="utf-8")
    with path.open("rb") as fh:
        up = client.post("/api/upload", files={"file": ("sample.txt", fh, "text/plain")})
    assert up.status_code == 200
    fid = up.json()["file_id"]
    st = client.post(f"/api/parse?file_id={fid}&wait=true").json()
    assert st["status"] == "FAILED"
    assert st["errors"][0]["code"] == "UNSUPPORTED_FORMAT"
