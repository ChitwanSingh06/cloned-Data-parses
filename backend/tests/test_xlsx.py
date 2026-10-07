from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.formats.xlsx_handler import parse_xlsx
from app.main import app
from app.models.block import BlockType

client = TestClient(app)


def _fixture(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws.append(["Region", "FY2025", "FY2026"])
    ws.append(["North", 1200, 1350])
    ws.append(["South", 980, 1010])
    ws.append(["East", 750, 820])
    tab = Table(displayName="RevenueTable", ref="A1:C4")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)

    ws2 = wb.create_sheet("Notes")
    ws2["A1"] = "Note"
    ws2["B1"] = "Value"
    ws2["A2"] = "Status"
    ws2["B2"] = "Approved"
    ws2["A3"] = "Updated"
    ws2["B3"] = "2026-10-07"

    out = tmp_path / "fixture.xlsx"
    wb.save(out)
    return out


def test_xlsx_parsing_canonical_output(tmp_path):
    path = _fixture(tmp_path)
    doc = parse_xlsx(str(path), document_id="xlsx12345678")

    assert doc.format == "xlsx"
    assert doc.status == "SUCCESS"
    assert doc.page_count == 2
    assert [p.page_number for p in doc.pages] == [1, 2]
    assert [b.content for b in doc.blocks if b.type == BlockType.heading] == ["Revenue", "Notes"]

    tables = [b for b in doc.blocks if b.type == BlockType.table]
    assert len(tables) == 2
    revenue = tables[0]
    assert revenue.content["n_rows"] == 4
    assert revenue.content["n_cols"] == 3
    assert revenue.content["headers"][0] == ["Region", "FY2025", "FY2026"]
    assert revenue.content["rows"][1] == ["South", 980, 1010]
    assert revenue.meta["worksheet"] == "Revenue"
    assert revenue.meta["range"] == "A1:C4"
    assert revenue.meta["excel_tables"] == [{"name": "RevenueTable", "ref": "A1:C4"}]
    assert revenue.meta["cell_provenance"][0]["cell"] == "A1"
    assert revenue.provenance and revenue.provenance.sources[0]["worksheet"] == "Revenue"
    assert revenue.bbox is None

    from app.core.config import PROCESSED_DIR
    out = PROCESSED_DIR / "xlsx12345678"
    assert (out / "document.json").exists()
    md = (out / "document.md").read_text(encoding="utf-8")
    assert "# Revenue" in md
    assert "| Region | FY2025 | FY2026 |" in md
    assert "| South | 980 | 1010 |" in md
    assert "# Notes" in md


def test_xlsx_upload_parse_downloads(tmp_path):
    path = _fixture(tmp_path)
    with path.open("rb") as fh:
        up = client.post(
            "/api/upload",
            files={"file": ("fixture.xlsx", fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert up.status_code == 200
    fid = up.json()["file_id"]
    st = client.post(f"/api/parse?file_id={fid}&wait=true").json()
    assert st["status"] == "SUCCESS"
    result = client.get(f"/api/result/{fid}").json()
    assert result["format"] == "xlsx"
    assert result["page_count"] == 2
    assert client.get(f"/api/download/{fid}/json").status_code == 200
    md = client.get(f"/api/download/{fid}/markdown")
    assert "Revenue" in md.text and "South" in md.text
    assert client.get(f"/api/documents/{fid}/page/1.png").status_code == 404


def test_corrupt_xlsx_is_structured(tmp_path):
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"PK\x03\x04not really an xlsx")
    doc = parse_xlsx(str(path), document_id="badxlsx123456")
    assert doc.status == "FAILED"
    assert doc.errors[0]["code"] in {"CORRUPT_XLSX", "PARSING_FAILED"}
